from openpyxl import Workbook

from rule_engine import RuleRegistry, WhitespaceCleansingRule


def test_whitespace_rule_detects_and_fixes_column_scoped_changes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Amount"])
    ws.append(["  pogi", 1])
    ws.append(["pogi\u00a0\u00a0pogi", 2])
    ws.append(["clean", 3])

    registry = RuleRegistry([WhitespaceCleansingRule()])
    results = registry.scan_sheet(ws, "Data", ["Name", "Amount"], start_row=2)
    proposals = registry.build_fix_proposals(results)

    assert len(results) == 2
    assert len(proposals) == 1
    assert proposals[0].column_name == "Name"
    assert proposals[0].affected_count == 2

    proposals[0].apply(wb)

    assert ws.cell(row=2, column=1).value == "pogi"
    assert ws.cell(row=3, column=1).value == "pogi pogi"
    assert ws.cell(row=4, column=1).value == "clean"
